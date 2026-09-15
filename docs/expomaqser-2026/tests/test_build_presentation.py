import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

import sys

SCRIPT_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import build_expomaqser_2026_presentation as pres  # noqa: E402


class PresentationBuilderTests(unittest.TestCase):
    def setUp(self):
        pres._FONT_PATHS_CACHE = None

    def test_build_assets_respects_refresh_flag(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_assets = Path(tmp_dir) / "assets"

            with patch.object(pres, "ASSETS_DIR", tmp_assets):
                calls = []

                def fake_draw(path, title, subtitle, motifs):
                    path.parent.mkdir(parents=True, exist_ok=True)
                    path.touch()
                    calls.append(path.name)

                with patch.object(pres, "draw_visual", side_effect=fake_draw):
                    first = pres.build_assets(refresh_assets=False)
                    self.assertEqual(len(calls), len(first))

                    pres.build_assets(refresh_assets=False)
                    self.assertEqual(len(calls), len(first), "Assets should not be regenerated when files exist")

                    pres.build_assets(refresh_assets=True)
                    self.assertEqual(len(calls), len(first) * 2, "Assets should regenerate when refresh is enabled")

    def test_verify_notes_persist_checks_exact_content(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            pptx_path = Path(tmp_dir) / "notes.pptx"
            deck = Presentation()
            deck.slide_width = pres.WIDE_W
            deck.slide_height = pres.WIDE_H

            expected = ["Nota de apertura", "Nota técnica"]
            for note in expected:
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                pres.add_note(slide, note)

            deck.save(pptx_path)
            pres.verify_notes_persist(pptx_path, expected)

            with self.assertRaises(RuntimeError):
                pres.verify_notes_persist(pptx_path, ["Nota equivocada", "Nota técnica"])

    def test_configure_font_paths_uses_explicit_fonts(self):
        with patch.object(pres.ImageFont, "truetype", return_value=object()) as mock_truetype:
            pres.configure_font_paths(
                font_regular=Path("/tmp/custom-regular.ttf"),
                font_bold=Path("/tmp/custom-bold.ttf"),
            )

        self.assertEqual(
            pres._FONT_PATHS_CACHE,
            ("/tmp/custom-regular.ttf", "/tmp/custom-bold.ttf"),
        )
        self.assertGreaterEqual(mock_truetype.call_count, 2)

    def test_configure_font_paths_raises_when_unresolvable(self):
        with patch.object(pres.ImageFont, "truetype", side_effect=OSError("missing font")):
            with self.assertRaises(RuntimeError):
                pres.configure_font_paths(
                    font_regular=Path("/tmp/not-found-regular.ttf"),
                    font_bold=Path("/tmp/not-found-bold.ttf"),
                )


if __name__ == "__main__":
    unittest.main()

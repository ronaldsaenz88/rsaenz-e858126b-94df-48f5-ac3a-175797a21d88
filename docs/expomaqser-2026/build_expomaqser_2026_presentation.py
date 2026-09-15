from __future__ import annotations

import argparse
from pathlib import Path
from typing import List, Tuple

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ---------- Paths ----------
BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "assets"
OUTPUT = BASE_DIR / "Mas_alla_de_la_maquinaria_EXPOMAQSER_2026.pptx"

# ---------- Deck constants ----------
WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

NAVY = RGBColor(19, 35, 52)
TEAL = RGBColor(0, 141, 150)
LIGHT_BG = RGBColor(246, 248, 250)
SOFT_GRAY = RGBColor(96, 108, 121)
DARK_TEXT = RGBColor(36, 44, 54)
WHITE = RGBColor(255, 255, 255)
ACCENT_ORANGE = RGBColor(228, 139, 61)
ALERT_RED = RGBColor(187, 67, 67)

PALETTE = {
    "navy": (26, 46, 66),
    "teal": (0, 141, 150),
    "orange": (228, 139, 61),
    "slate": (89, 106, 124),
    "light": (240, 244, 247),
}


_FONT_PATHS_CACHE: tuple[str, str] | None = None


def _resolve_font_paths(font_regular: Path | None = None, font_bold: Path | None = None) -> tuple[str, str]:
    """Resolve deterministic font files required to generate reproducible PNG assets."""
    regular_candidates = [
        str(font_regular) if font_regular else "",
        "DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/DejaVuSans.ttf",
        "C:/Windows/Fonts/DejaVuSans.ttf",
    ]
    bold_candidates = [
        str(font_bold) if font_bold else "",
        "DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/Library/Fonts/DejaVuSans-Bold.ttf",
        "C:/Windows/Fonts/DejaVuSans-Bold.ttf",
    ]

    def _first_loadable(candidates: List[str]) -> str | None:
        for path in candidates:
            if not path:
                continue
            try:
                ImageFont.truetype(path, size=12)
                return path
            except OSError:
                continue
        return None

    regular_path = _first_loadable(regular_candidates)
    bold_path = _first_loadable(bold_candidates)

    if not regular_path or not bold_path:
        raise RuntimeError(
            "DejaVuSans regular/bold fonts are required to generate reproducible visual assets. "
            "Install DejaVu fonts or provide these files before running the script."
        )

    return regular_path, bold_path


def configure_font_paths(font_regular: Path | None = None, font_bold: Path | None = None) -> None:
    global _FONT_PATHS_CACHE
    _FONT_PATHS_CACHE = _resolve_font_paths(font_regular=font_regular, font_bold=font_bold)


def _font(size: int, bold: bool = False):
    global _FONT_PATHS_CACHE
    if _FONT_PATHS_CACHE is None:
        _FONT_PATHS_CACHE = _resolve_font_paths()
    regular_path, bold_path = _FONT_PATHS_CACHE
    font_path = bold_path if bold else regular_path
    return ImageFont.truetype(font_path, size=size)


def draw_visual(path: Path, title: str, subtitle: str, motifs: List[str]) -> None:
    """Create simple, license-safe custom visual blocks for the deck."""
    img = Image.new("RGB", (1600, 900), PALETTE["light"])
    d = ImageDraw.Draw(img)

    # background bands
    d.rectangle([0, 0, 1600, 140], fill=PALETTE["navy"])
    d.rectangle([0, 140, 1600, 170], fill=PALETTE["teal"])

    # motif cards
    card_w = 460
    x = 70
    y = 250
    for idx, motif in enumerate(motifs[:3]):
        fill = PALETTE["teal"] if idx % 2 == 0 else PALETTE["orange"]
        d.rounded_rectangle([x, y, x + card_w, y + 360], radius=18, fill=(255, 255, 255), outline=fill, width=6)
        d.rectangle([x + 28, y + 28, x + 80, y + 80], fill=fill)
        d.text((x + 100, y + 30), motif, font=_font(34, bold=True), fill=PALETTE["navy"])
        d.text(
            (x + 36, y + 120),
            "Aplicación en obra\ncon foco en control,\nproductividad y decisiones.",
            font=_font(28),
            fill=PALETTE["slate"],
            spacing=8,
        )
        x += card_w + 50

    d.text((70, 42), title, font=_font(50, bold=True), fill=(255, 255, 255))
    d.text((70, 182), subtitle, font=_font(36), fill=PALETTE["navy"])

    img.save(path)


def build_assets(
    refresh_assets: bool = False,
    font_regular: Path | None = None,
    font_bold: Path | None = None,
) -> dict[str, Path]:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    configure_font_paths(font_regular=font_regular, font_bold=font_bold)
    visuals: dict[str, Tuple[str, str, List[str]]] = {
        "opening": (
            "EXPOMAQSER 2026 | Guayaquil",
            "Construcción y real estate: del dato a la decisión",
            ["Retrasos", "Sobrecostos", "Coordinación"],
        ),
        "context": (
            "Contexto local",
            "Proyectos urbanos y presión por plazos",
            ["Demanda", "Financiamiento", "Riesgo operativo"],
        ),
        "stack": (
            "Capa digital",
            "BIM + campo + negocio + IA",
            ["Integración", "Trazabilidad", "Gobernanza"],
        ),
        "students": (
            "Talento Civil + Sistemas",
            "Nuevos perfiles para la obra digital",
            ["Datos", "Automatización", "Liderazgo"],
        ),
        "case": (
            "Caso Torre Santiago",
            "Simulación realista para decisión gerencial",
            ["Cronograma", "Productividad", "Caja"],
        ),
        "cv": (
            "Computer Vision en campo",
            "Seguridad, avance y control visual",
            ["EPP", "Frentes", "Alertas"],
        ),
        "close": (
            "Siguiente paso",
            "Pilotos rápidos con impacto medible",
            ["ROI", "Escalamiento", "Cultura"],
        ),
    }

    out: dict[str, Path] = {}
    for key, (title, subtitle, motifs) in visuals.items():
        img_path = ASSETS_DIR / f"{key}.png"
        if refresh_assets or not img_path.exists():
            draw_visual(img_path, title, subtitle, motifs)
        out[key] = img_path
    return out


def add_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, title: str, subtitle: str | None = None, dark=False):
    color = WHITE if dark else NAVY
    tx = slide.shapes.add_textbox(Inches(0.65), Inches(0.38), Inches(8.4), Inches(1.2))
    p = tx.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = color

    if subtitle:
        st = slide.shapes.add_textbox(Inches(0.68), Inches(1.08), Inches(8.5), Inches(0.8))
        p = st.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(14)
        p.font.color.rgb = WHITE if dark else SOFT_GRAY


def add_footer(slide, n: int, dark=False):
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(7.05), Inches(12.05), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()

    tx = slide.shapes.add_textbox(Inches(11.7), Inches(7.1), Inches(0.8), Inches(0.3))
    p = tx.text_frame.paragraphs[0]
    p.text = f"{n:02d}"
    p.font.size = Pt(9)
    p.font.color.rgb = WHITE if dark else SOFT_GRAY
    p.alignment = PP_ALIGN.RIGHT


def add_bullets(slide, items: List[str], x=0.78, y=1.75, w=7.25, h=4.9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(21 if i == 0 else 19)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(12)


def add_visual(slide, image_path: Path):
    slide.shapes.add_picture(str(image_path), Inches(8.25), Inches(1.45), Inches(4.55), Inches(5.25))


def add_note(slide, text: str):
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.paragraphs[0].text = text


def verify_notes_persist(path: Path, expected_notes: List[str]) -> None:
    """Ensure generated speaker notes survive save/reload in the output deck."""
    reloaded = Presentation(path)
    reloaded_notes = [slide.notes_slide.notes_text_frame.text.strip() for slide in reloaded.slides]

    if len(reloaded_notes) != len(expected_notes):
        raise RuntimeError(
            f"Speaker notes persistence check failed: expected {len(expected_notes)} notes, found {len(reloaded_notes)}."
        )

    for idx, (expected, actual) in enumerate(zip(expected_notes, reloaded_notes), start=1):
        if actual != expected.strip():
            raise RuntimeError(
                "Speaker notes persistence check failed "
                f"on slide {idx}: expected '{expected.strip()}', found '{actual}'."
            )


def add_title_only_visual(prs: Presentation, n: int, title: str, subtitle: str, image_path: Path, note: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, NAVY)
    add_header(s, title, subtitle, dark=True)
    s.shapes.add_picture(str(image_path), Inches(0.65), Inches(1.8), Inches(12.05), Inches(4.8))
    add_footer(s, n, dark=True)
    add_note(s, note)


def build_deck(
    refresh_assets: bool = False,
    output_path: Path = OUTPUT,
    font_regular: Path | None = None,
    font_bold: Path | None = None,
):
    """Build and save the EXPOMAQSER deck, returning (output_path, slide_count).

    Raises RuntimeError if required fonts cannot be resolved or if speaker notes
    fail the save/reload persistence check.
    """
    visuals = build_assets(
        refresh_assets=refresh_assets,
        font_regular=font_regular,
        font_bold=font_bold,
    )

    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H

    # 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, NAVY)
    add_header(
        s,
        "Más allá de la maquinaria",
        "EXPOMAQSER 2026 | Guayaquil | Estudiantes + sector constructor e inmobiliario",
        dark=True,
    )
    s.shapes.add_textbox(Inches(0.8), Inches(2.15), Inches(6.8), Inches(1.2)).text_frame.text = (
        "La ventaja competitiva en 2026 no está solo en equipos:"
    )
    p = s.shapes[-1].text_frame.paragraphs[0]
    p.font.size = Pt(23)
    p.font.color.rgb = WHITE
    p.font.bold = True

    s.shapes.add_textbox(Inches(0.8), Inches(3.25), Inches(6.9), Inches(2.1)).text_frame.text = (
        "está en cómo conectamos campo, datos, software e IA para tomar mejores decisiones."
    )
    p = s.shapes[-1].text_frame.paragraphs[0]
    p.font.size = Pt(28)
    p.font.color.rgb = TEAL
    p.font.bold = True

    s.shapes.add_picture(str(visuals["opening"]), Inches(7.8), Inches(1.6), Inches(4.5), Inches(4.8))
    add_footer(s, 1, dark=True)
    add_note(
        s,
        "Abrir con una tesis clara para ambos públicos: la obra necesita control operativo y lectura digital. "
        "Conectar expectativas del empresario (ROI) y del estudiante (empleabilidad).",
    )

    # 2
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Ruta de la charla", "30 min charla + 10 min demo + 5 min preguntas")
    add_bullets(
        s,
        [
            "Bloque 1 (10 min): desafíos reales de construcción y real estate en Guayaquil.",
            "Bloque 2 (10 min): capa digital, datos operativos e IA aplicada.",
            "Bloque 3 (10 min): caso Torre Santiago, decisiones y plan de adopción.",
            "Demo (10 min): tablero + alerta predictiva + recomendación accionable.",
            "Q&A (5 min): cómo arrancar en universidad y en empresa este año.",
        ],
    )
    add_visual(s, visuals["context"])
    add_footer(s, 2)
    add_note(s, "Marcar tiempos desde el inicio genera expectativa y orden. Confirmar que habrá demo real al final.")

    # 3
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Guayaquil hoy: presión por ejecutar mejor", "Más demanda, más escrutinio, menos margen para improvisar")
    add_bullets(
        s,
        [
            "El mercado exige plazos confiables para proyectos residenciales, mixtos y logísticos.",
            "Los sobrecostos impactan rentabilidad del desarrollador y flujo del contratista.",
            "Proveedores necesitan previsibilidad de compras para no romper cadena de suministro.",
            "La coordinación entre diseño, obra y comercial sigue siendo un cuello de botella.",
            "Resultado: quien controla datos y decisiones, controla el negocio.",
        ],
    )
    add_visual(s, visuals["context"])
    add_footer(s, 3)
    add_note(s, "Aterrizar contexto local: no hablar de IA abstracta, hablar de presión real de caja, plazo y reputación.")

    # 4
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "El costo oculto de la fragmentación", "Tenemos datos, pero dispersos en múltiples silos")
    add_bullets(
        s,
        [
            "Cronograma vive en un archivo; avance real, en fotos y chats.",
            "Costos y compras se analizan tarde, cuando el desvío ya ocurrió.",
            "Riesgos de seguridad se reportan sin trazabilidad comparativa.",
            "Dirección recibe reportes semanales, no señales tempranas diarias.",
            "Sin integración, reaccionamos tarde y pagamos más.",
        ],
    )
    add_visual(s, visuals["stack"])
    add_footer(s, 4)
    add_note(s, "Conectar con dolor operativo del sector: retraso detectado tarde = costo financiero y contractual.")

    # 5
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Qué llamamos ‘capa digital’", "Un sistema para pasar de dato suelto a decisión de obra")
    add_bullets(
        s,
        [
            "Captura: maquinaria, partes diarios, BIM, clima, fotos de campo.",
            "Integración: un modelo de datos común por frente, actividad y costo.",
            "Analítica: indicadores de productividad, seguridad, avance y desvío.",
            "IA: alertas predictivas y explicación de causas probables.",
            "Acción: reunión de control con decisiones trazables y responsables.",
        ],
    )
    add_visual(s, visuals["stack"])
    add_footer(s, 5)
    add_note(s, "Definir capa digital como proceso de gestión, no como software aislado.")

    # 6
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Valor para industria: ROI operativo", "Desarrolladores, contratistas y proveedores")
    add_bullets(
        s,
        [
            "Menos retrasos críticos por alertas tempranas en ruta de obra.",
            "Mayor productividad al detectar frentes subutilizados.",
            "Menos reprocesos por mejor coordinación diseño-campo.",
            "Mejor seguridad con monitoreo de comportamientos de riesgo.",
            "Más confianza para inversionistas por trazabilidad y control.",
        ],
    )
    add_visual(s, visuals["opening"])
    add_footer(s, 6)
    add_note(s, "Hablar en lenguaje de negocio: riesgo, productividad, cumplimiento y reputación.")

    # 7
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Valor para estudiantes: empleabilidad real", "Perfil híbrido Civil + Sistemas como ventaja profesional")
    add_bullets(
        s,
        [
            "Comprender obra + modelar datos = perfil altamente demandado.",
            "Nuevos roles: planner analítico, coordinador BIM-datos, PM digital.",
            "Competencias clave: SQL, dashboards, lectura de KPIs de obra, IA aplicada.",
            "La práctica profesional migrará de reporte manual a gestión asistida por datos.",
            "No reemplazo: mayor capacidad de criterio técnico y liderazgo.",
        ],
    )
    add_visual(s, visuals["students"])
    add_footer(s, 7)
    add_note(s, "Mensaje motivacional con realismo: la tecnología amplía el rol del ingeniero, no lo elimina.")

    # 8
    add_title_only_visual(
        prs,
        8,
        "Caso de estudio: Torre Santiago (ficticio, plausible)",
        "Edificio residencial de 22 niveles en Guayaquil norte",
        visuals["case"],
        "Presentar el caso como simulación práctica para el demo: datos razonables y decisiones tipo comité de obra.",
    )

    # 9
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Torre Santiago: línea base", "Semana 24 de un plan de 14 meses")
    add_bullets(
        s,
        [
            "Presupuesto contractual: USD 14.8 millones.",
            "Avance planificado: 58% | avance real: 49%.",
            "Equipos críticos: grúa torre, bomba de hormigón, encofrado.",
            "Incidentes de seguridad leves: 6 (meta trimestral: <=2).",
            "Riesgo de caja: certificación del hito comercial en 5 semanas.",
        ],
    )
    add_visual(s, visuals["case"])
    add_footer(s, 9)
    add_note(s, "Dar credibilidad financiera y operativa. Mostrar por qué un retraso técnico se vuelve riesgo comercial.")

    # 10
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Qué revela la capa de datos", "Cruce de productividad, clima y secuencia crítica")
    add_bullets(
        s,
        [
            "Frente de estructura con productividad -11% por interferencias logísticas.",
            "Disponibilidad de bomba de hormigón cayó a 78% durante 3 semanas.",
            "Compras de acero con variabilidad de entrega (+5 días promedio).",
            "Retrabajos en losas por coordinación tardía de instalaciones.",
            "Conclusión: el retraso no es una causa única, sino un patrón sistémico.",
        ],
    )
    add_visual(s, visuals["stack"])
    add_footer(s, 10)
    add_note(s, "Enseñar lectura causal: cómo varias señales pequeñas producen un desvío mayor.")

    # 11
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "IA predictiva: alerta temprana accionable", "Modelo estima probabilidad de incumplir hito de semana 30")
    add_bullets(
        s,
        [
            "Riesgo estimado de atraso del hito: 76% (señal alta).",
            "Factores explicativos: productividad, disponibilidad de equipo y secuencia crítica.",
            "Ventana de intervención: próximas 2 semanas.",
            "Escenario recomendado: reforzar cuadrilla + reprogramar vaciados nocturnos.",
            "Impacto esperado: recuperar 4 a 6 puntos de avance en 6 semanas.",
        ],
    )
    add_visual(s, visuals["close"])
    add_footer(s, 11)
    add_note(s, "Aclarar que la IA no reemplaza la planificación: prioriza dónde actuar primero.")

    # 12
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Computer Vision en operación", "Del video de obra a alertas de seguridad y progreso")
    add_bullets(
        s,
        [
            "Detección de uso de EPP en zonas definidas de riesgo.",
            "Conteo de avance visual por elemento construido vs. BIM.",
            "Identificación de zonas congestionadas para optimizar logística.",
            "Reporte automático diario para jefatura de obra y HSE.",
            "Beneficio: supervisión más consistente sin depender solo del recorrido manual.",
        ],
    )
    add_visual(s, visuals["cv"])
    add_footer(s, 12)
    add_note(s, "Mantener promesa realista: CV ayuda a detectar y priorizar, la validación final es humana.")

    # 13
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Decisión integrada: comité semanal con evidencia", "Civil + Sistemas + Finanzas + Operaciones")
    add_bullets(
        s,
        [
            "Dato común: una sola versión de cronograma, costo y avance.",
            "Priorización por impacto económico y riesgo contractual.",
            "Asignación clara de responsables y fechas de cierre.",
            "Seguimiento de cumplimiento con tablero ejecutivo.",
            "Resultado: menos discusión subjetiva, más ejecución coordinada.",
        ],
    )
    add_visual(s, visuals["students"])
    add_footer(s, 13)
    add_note(s, "Subrayar colaboración multidisciplinaria: el valor aparece cuando áreas comparten el mismo dato.")

    # 14
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, LIGHT_BG)
    add_header(s, "Plan de adopción en 90 días", "Piloto realista para empresa o alianza universidad-industria")
    add_bullets(
        s,
        [
            "Semana 1-2: definir caso prioritario (retraso, seguridad o productividad).",
            "Semana 3-5: conectar fuentes mínimas de datos y tablero operativo.",
            "Semana 6-8: correr modelo predictivo y validar con equipo de obra.",
            "Semana 9-12: medir ROI (tiempo, costo, incidentes, retrabajo).",
            "Escalamiento: estandarizar metodología para nuevos proyectos.",
        ],
    )
    add_visual(s, visuals["close"])
    add_footer(s, 14)
    add_note(s, "Cerrar con acción concreta: empezar pequeño, medir impacto, luego escalar.")

    # 15
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, NAVY)
    add_header(
        s,
        "Cierre",
        "La próxima ventaja competitiva en construcción ecuatoriana será técnico-digital",
        dark=True,
    )
    quote = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(7.2), Inches(2.6))
    q = quote.text_frame.paragraphs[0]
    q.text = "No se trata de tener más datos."
    q.font.size = Pt(36)
    q.font.bold = True
    q.font.color.rgb = WHITE
    q = quote.text_frame.add_paragraph()
    q.text = "Se trata de convertirlos en decisiones a tiempo."
    q.font.size = Pt(34)
    q.font.bold = True
    q.font.color.rgb = TEAL

    s.shapes.add_picture(str(visuals["opening"]), Inches(8.05), Inches(1.7), Inches(4.2), Inches(4.7))

    c = s.shapes.add_textbox(Inches(0.8), Inches(5.25), Inches(7.0), Inches(1.4))
    cp = c.text_frame.paragraphs[0]
    cp.text = "Siguiente: demo en vivo (10 min) + preguntas (5 min)."
    cp.font.size = Pt(20)
    cp.font.color.rgb = WHITE

    add_footer(s, 15, dark=True)
    add_note(
        s,
        "Invitar al público a aterrizar un caso propio: qué decisión crítica mejorarían mañana si integran datos de campo y negocio.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    expected_notes = [slide.notes_slide.notes_text_frame.text.strip() for slide in prs.slides]
    verify_notes_persist(output_path, expected_notes=expected_notes)
    return output_path, len(prs.slides)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build revised EXPOMAQSER 2026 presentation.")
    parser.add_argument(
        "--refresh-assets",
        action="store_true",
        help="Regenerate PNG visual assets even if they already exist.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=OUTPUT,
        help="Output .pptx path. Defaults to docs/expomaqser-2026/Mas_alla_de_la_maquinaria_EXPOMAQSER_2026.pptx",
    )
    parser.add_argument(
        "--font-regular",
        type=Path,
        default=None,
        help="Optional path to the regular DejaVuSans font file for deterministic asset generation.",
    )
    parser.add_argument(
        "--font-bold",
        type=Path,
        default=None,
        help="Optional path to the bold DejaVuSans font file for deterministic asset generation.",
    )
    args = parser.parse_args()
    saved_path, slide_count = build_deck(
        refresh_assets=args.refresh_assets,
        output_path=args.output,
        font_regular=args.font_regular,
        font_bold=args.font_bold,
    )
    print(f"Presentation generated: {saved_path}")
    print(f"Slides: {slide_count}")
